import docx
import xml.etree.ElementTree as ET

doc = docx.Document('devops report.docx')
rel_map = {rId: rel.target_ref for rId, rel in doc.part.rels.items() if 'image' in rel.target_ref}

for i, p in enumerate(doc.paragraphs):
    if 'drawing' in p._element.xml:
        root = ET.fromstring(p._element.xml)
        for elem in root.iter():
            if elem.tag.endswith('blip'):
                embed = elem.attrib.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed')
                target = rel_map.get(embed, 'unknown')
                caption = doc.paragraphs[i+1].text if i+1 < len(doc.paragraphs) else ''
                prev_text = doc.paragraphs[i-1].text if i-1 >= 0 else ''
                caption_clean = caption.encode('ascii', 'ignore').decode().strip()
                prev_clean = prev_text.encode('ascii', 'ignore').decode().strip()
                print(f"P{i}: {target} -> Above: '{prev_clean[:40]}' | Caption: '{caption_clean[:60]}'")
