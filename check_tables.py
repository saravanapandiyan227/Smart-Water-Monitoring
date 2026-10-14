import docx

doc = docx.Document('devops report.docx')
print('Tables count:', len(doc.tables))
for i, t in enumerate(doc.tables):
    rows = len(t.rows)
    cols = len(t.columns)
    first_cell = t.cell(0, 0).text.strip() if rows > 0 and cols > 0 else ''
    last_cell = t.cell(rows-1, cols-1).text.strip() if rows > 0 and cols > 0 else ''
    print(f"Table {i}: {rows} rows x {cols} cols")
    print(f"   [0,0]: '{first_cell[:50]}'")
    print(f"   [-1,-1]: '{last_cell[:50]}'")
