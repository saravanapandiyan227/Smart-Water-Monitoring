import docx

doc = docx.Document('devops report.docx')

img_pars = [12, 110, 138, 140, 159, 182, 207, 234, 237, 239, 255, 256, 279, 304, 305, 322, 324, 338, 341, 357, 380]

for p_idx in img_pars:
    print(f"\n=== Context for Image at Paragraph {p_idx} ===")
    start = max(0, p_idx - 2)
    end = min(len(doc.paragraphs), p_idx + 4)
    for j in range(start, end):
        txt = doc.paragraphs[j].text.encode('ascii', 'ignore').decode().strip()
        mark = "--> [IMAGE]" if j == p_idx else "   "
        print(f"{mark} P{j}: {txt}")
